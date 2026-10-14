import tkinter as tk
from tkinter import messagebox, filedialog
import re
from datetime import datetime, timedelta
import traceback
import os
import webbrowser
import calendar
import importlib.util
from pathlib import Path
import sys
from dataclasses import dataclass
from typing import Callable, List, Optional, Tuple


# ==========================================
# 0. STARTUP DEPENDENCY CHECK
# ==========================================
REQUIRED_DEPENDENCIES = {
    "customtkinter": "customtkinter",
    "packaging": "packaging",
    "openpyxl": "openpyxl",
    "beautifulsoup4": "bs4",
}


def check_required_dependencies():
    """Stop cleanly with installation instructions if a required package is missing."""
    missing = [package for package, module in REQUIRED_DEPENDENCIES.items()
               if importlib.util.find_spec(module) is None]
    if not missing:
        return

    requirements_file = Path(__file__).with_name("requirements.txt")
    if requirements_file.exists():
        install_command = f'"{sys.executable}" -m pip install -r "{requirements_file}"'
    else:
        install_command = f'"{sys.executable}" -m pip install ' + " ".join(missing)

    message = (
        "Spreader Creator cannot start because required Python packages are missing:\n\n"
        f"{', '.join(missing)}\n\n"
        "Install them from Command Prompt with:\n\n"
        f"{install_command}\n\n"
        "Then launch Spreader Creator again."
    )
    try:
        dialog_root = tk.Tk()
        dialog_root.withdraw()
        messagebox.showerror("Missing Dependencies", message)
        dialog_root.destroy()
    except tk.TclError:
        # Keep a useful message if Tk itself cannot create a window.
        print(message, file=sys.stderr)
    raise SystemExit(1)


check_required_dependencies()

import customtkinter as ctk
import openpyxl
from bs4 import BeautifulSoup

# ==========================================
# 1. CONFIGURATION & THEMES
# ==========================================
TEMPLATE_PATH = r"\\ITPAATPNFP001.ad.faa.gov\TPA-Data\Air Traffic\Watch Desk Documents\DAILY SPREADER -  2026.xlsm"
OUTPUT_DIR = r"\\ITPAATPNFP001.ad.faa.gov\TPA-Data\Air Traffic\Watch Desk Documents\\"
DOWNLOADS_DIR = r"\\ITPAATPNFP001.ad.faa.gov\TPA-Data\Air Traffic\Watch Desk Documents\Downloads"
WMT_URL = "https://wmtscheduler.faa.gov/Views/WorksheetView"
APP_ICON = r"\\ITPAATPNFP001.ad.faa.gov\TPA-Data\Air Traffic\Watch Desk Documents\IC-Tools\Spreader Creator\Spreader_Creator_Icon.ico"

CONFIG = {
    # Fine-tune the complete visual theme here.
    # Dark navy / blue / gold palette inspired by the latest reference.
    "bg_main": "#101C2B",
    "bg_surface": "#1A2F4C",
    "bg_surface_alt": "#243B55",
    "bg_input": "#14263B",
    "bg_log": "#0B1522",
    "fg_main": "#F0F4F8",
    "fg_muted": "#B8C7D9",
    "fg_log": "#D8E8F5",
    "fg_on_accent": "#1A2F4C",
    # Gold is intentionally reserved for primary actions and highlights.
    "accent": "#D99B26",
    "accent_hover": "#E5AF3C",
    "accent_pressed": "#B98017",
    "btn_secondary": "#3C619D",
    "btn_sec_hover": "#5278B5",
    "border": "#35516F",
    "success": "#63C7AA",
    "danger": "#FF8FA8",
    "radius_card": 24,
    "radius_control": 14,
    "font_family": "Segoe UI"
}

# Backwards-compatible alias for the non-UI code and any local customizations.
COLORS = CONFIG


@dataclass
class ParsedData:
    """Records extracted from one WMT worksheet export."""
    shifts: List[dict]
    leave_list: List[str]
    notes_list: List[str]


def validate_date(date_str: str) -> str:
    """Validate and normalize the date used in the workbook filename."""
    parsed = datetime.strptime(date_str.strip(), "%m-%d-%Y")
    return f"{parsed.month}-{parsed.day:02d}-{parsed.year}"

# ==========================================
# 2. CORE BUSINESS LOGIC & TIME MATH
# ==========================================
def to_num(val: str):
    if not val:
        return ""
    return int(val) if val.isdigit() else val

def subtract_hours(time_str: str, hours_to_sub: int) -> str:
    if time_str in ["MID", "2400"]:
        time_str = "0000"
    try:
        total_minutes = int(time_str[:2]) * 60 + int(time_str[2:4])
        total_minutes = (total_minutes - hours_to_sub * 60) % (24 * 60)
        return f"{total_minutes // 60:02d}{total_minutes % 60:02d}"
    except (TypeError, ValueError):
        return time_str

def add_hours(h: int, m: int, hours_to_add: int) -> str:
    total_minutes = (h * 60 + m + hours_to_add * 60) % (24 * 60)
    if total_minutes == 0 and hours_to_add > 0:
        return "2400"
    return f"{total_minutes // 60:02d}{total_minutes % 60:02d}"

def calculate_times(start_time: str, final_code: str, orig_code: str,
                    ot_hours: int = 0, fos_hours: int = 0) -> Tuple[str, str]:
    if start_time in ["2230", "2300", "2330", "2400"]:
        return ("", "MID") if final_code == "T" else ("MID", "")

    # FOS Start time adjustment for calculation purposes only
    adjusted_start = subtract_hours(start_time, fos_hours)

    try:
        start_h, start_m = int(adjusted_start[:2]), int(adjusted_start[2:4])
    except (TypeError, ValueError):
        return "", ""

    if final_code == "T":       # Dev shift
        shift_len = ot_hours if ot_hours >= 8 else 8 + fos_hours
        return "", add_hours(start_h, start_m, shift_len)
    elif final_code == "S":     # Standard
        return add_hours(start_h, start_m, 7), ""
    elif final_code == "N":     # 9-hour
        return add_hours(start_h, start_m, 8), add_hours(start_h, start_m, 9)
    elif final_code == "A":     # 10-hour
        return add_hours(start_h, start_m, 8), add_hours(start_h, start_m, 10)
    elif final_code == "$$":    # Extended OT
        if orig_code == "A" or ot_hours == 0:
            return add_hours(start_h, start_m, 8), add_hours(start_h, start_m, 10)
        elif ot_hours > 8:
            return add_hours(start_h, start_m, 8), add_hours(start_h, start_m, ot_hours)
        else:
            return add_hours(start_h, start_m, 8), add_hours(start_h, start_m, 8 + ot_hours)
    elif final_code == "$":     # Standard OT
        return add_hours(start_h, start_m, 8), ""
    else:
        return add_hours(start_h, start_m, 8), ""

def sort_chronological(shift: dict) -> int:
    t = int(shift['start'])
    return t - 2400 if t >= 2200 else t


# ==========================================
# 3. DATA EXTRACTION (HTML PARSING)
# ==========================================
def parse_html_data(filepath: str, log_cb: Callable[[str], None]) -> ParsedData:
    filepath = filepath.strip('"').strip("'")
    log_cb(f"Opening file: {filepath}")

    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        soup = BeautifulSoup(f.read(), 'html.parser')

    # --- 1. OVERTIME ---
    ot_dict = {}
    log_cb("--- SCANNING OVERTIME TABLE ---")
    for row in soup.find_all('tr'):
        cols = row.find_all('td')
        if len(cols) == 5 and cols[1].get_text(strip=True).casefold() in {"scheduled", "call in", "holdover"}:
            init = cols[0].get_text(strip=True).upper()
            hrs_text = cols[2].get_text(strip=True)
            reason_text = cols[3].get_text(strip=True).upper()

            if hrs_text.isdigit():
                hrs = int(hrs_text)
                fos_match = re.search(r'(\d+)\s*FOS', reason_text)
                fos = int(fos_match.group(1)) if fos_match else 0

                if init not in ot_dict: ot_dict[init] = {'hrs': 0, 'fos': 0}
                ot_dict[init]['hrs'] += hrs
                ot_dict[init]['fos'] += fos
                log_cb(f"OT Captured -> {init}: {hrs} hr(s) | {fos} FOS")

    # --- 2. REQUESTS ---
    leave_list = []
    log_cb("--- SCANNING LEAVE/REQUESTS ---")
    for table_id in ['tblShiftChange', 'tblXTECTE', 'tblSupRequests']:
        req_table = soup.find('table', id=table_id)
        if not req_table: continue

        for row in req_table.find_all('tr'):
            cols = row.find_all('td')
            if len(cols) >= 6:
                init, lv_type, lv_from, lv_to = [cols[i].get_text(strip=True) for i in range(4)]
                status = cols[5].get_text(strip=True)

                if ("approved" not in status.casefold()
                        and "shift chg" not in lv_type.casefold()
                        and "shift swap" not in lv_type.casefold()):
                    time_addon = ""
                    if lv_from and lv_from.lower() != "shift":
                        time_addon = f" ({lv_from} - {lv_to})" if lv_to else f" ({lv_from})"
                    leave_entry = f"{init} - {lv_type}{time_addon}"
                    leave_list.append(leave_entry)
                    log_cb(f"Flagged -> {leave_entry} (Status: {status})")

    # --- 3. NOTES ---
    notes_list = []
    log_cb("--- SCANNING SHIFT NOTES ---")
    notes_table = soup.find('table', id='ShiftNotesForDay')
    if notes_table:
        for row in notes_table.find_all('tr'):
            cols = row.find_all('td')
            if len(cols) >= 2 and cols[1].find('a', title=True):
                note_text = cols[1].find('a', title=True).get_text(strip=True)
                if note_text and note_text != "<Add Shift Notes>":
                    cleaned_note = re.sub(r'^[A-Za-z]{2,3}\s*-\s*', '', note_text)

                    # Filter out WMT auto-generated L&A/CruArt notes (e.g., "65 - OT..." or "51 - LOT...")
                    if re.match(r'^\d{2}\s*-', cleaned_note):
                        log_cb(f"Ignored System Note -> {cleaned_note}")
                    else:
                        notes_list.append(cleaned_note)
                        log_cb(f"Note -> {cleaned_note}")

    # --- 4. SHIFTS ---
    shifts = []
    log_cb("--- PROCESSING SHIFTS ---")
    ignore_tables = ['tblShiftChange', 'tblXTECTE', 'tblSupRequests', 'ShiftNotesForDay']

    for row in soup.find_all('tr'):
        cols = row.find_all('td')
        if len(cols) < 4: continue

        parent = row.find_parent('table')
        if parent and parent.get('id') in ignore_tables: continue

        time_str = cols[0].get_text(strip=True).replace('\xa0', '').strip()
        time_match = re.match(r'^(\d{2,4})([A-Z]*)$', time_str)

        if time_match:
            t_str, orig_code = time_match.groups()
            t_str = t_str.zfill(4) if len(t_str) == 3 else (t_str + "00" if len(t_str) <= 2 else t_str)

            for idx, col in enumerate(cols[2:4]):
                is_dev = (idx == 1)
                cell_words = col.get_text(separator=' ', strip=True).replace('\xa0', ' ').split()

                for word in cell_words:
                    match = re.match(r'^([A-Z]{2})(\$?)$', word)
                    if match:
                        init, has_dollar = match.group(1), bool(match.group(2))
                        if init in ["OF", "OT"]: continue

                        ot_hrs = ot_dict.get(init, {}).get('hrs', 0)
                        fos_hrs = ot_dict.get(init, {}).get('fos', 0)

                        # Determine Code
                        if is_dev:
                            final_code = "T"
                        else:
                            final_code = orig_code
                            if has_dollar:
                                final_code = "$$" if orig_code == "A" or ot_hrs > 8 else "$"
                            elif not orig_code and (0 < ot_hrs < 8) and not has_dollar:
                                final_code = "$$"

                        if fos_hrs > 0:
                            log_cb(f"FOS Detected -> {init} stays in {t_str} block, off-time math runs from {subtract_hours(t_str, fos_hrs)}")

                        shifts.append({
                            'initials': init,
                            'start': t_str,
                            'orig_code': orig_code,
                            'code': final_code,
                            'ot': ot_hrs,
                            'fos': fos_hrs
                        })

                        log_cb(f"Parsed -> {init} | Start: {t_str} | Code: {final_code or 'Standard (8-hr)'}{' (DEV)' if is_dev else ''}")

    return ParsedData(shifts, leave_list, notes_list)


# ==========================================
# 4. EXCEL GENERATION
# ==========================================
def build_excel_file(all_shifts: List[dict], leave_list: List[str], notes_list: List[str],
                     date_str: str, log_cb: Callable[[str], None],
                     dry_run: bool = False) -> Optional[str]:
    safe_date = validate_date(date_str)
    if dry_run:
        log_cb("DRY RUN: input validated; no workbook was opened or saved.")
        return None

    morning_shifts = [s for s in all_shifts if s['start'] in ["2230", "2300", "2330", "2400"] or int(s['start']) < 1130]
    afternoon_shifts = [s for s in all_shifts if s not in morning_shifts]

    morning_shifts.sort(key=sort_chronological)
    afternoon_shifts.sort(key=sort_chronological)

    log_cb("Loading Excel template (This may take a moment)...")
    template_path = Path(TEMPLATE_PATH)
    output_dir = Path(OUTPUT_DIR)
    if not template_path.is_file():
        raise FileNotFoundError(f"Excel template not found: {template_path}")
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(template_path, keep_vba=True)
    if "Projected Coverage" not in wb.sheetnames:
        raise KeyError("Template is missing the 'Projected Coverage' worksheet")
    ws = wb["Projected Coverage"]

    # Clear writable areas so records from a previous, larger run cannot remain.
    for row in range(7, 109):
        for col in ("E", "F", "G", "H", "J", "K", "L", "M"):
            ws[f"{col}{row}"] = None
    for row in list(range(9, 19)) + list(range(19, 29)):
        ws[f"AG{row}"] = None

    def write_block(shifts_to_write, cols, label, start_row):
        row = start_row
        last_start = ""
        c_init, c_off, c_code, c_req = cols

        for shift in shifts_to_write:
            if last_start and shift['start'] != last_start:
                row += 1

            if row >= start_row + 100:
                raise ValueError("Too many shifts for the available template rows")

            off, req = calculate_times(shift['start'], shift['code'], shift['orig_code'], shift['ot'], shift['fos'])
            ws[f'{c_init}{row}'] = shift['initials']
            ws[f'{c_off}{row}'] = to_num(off)
            ws[f'{c_code}{row}'] = shift['code']
            ws[f'{c_req}{row}'] = to_num(req)

            last_start = shift['start']
            row += 1
        log_cb(f"Wrote {len(shifts_to_write)} {label} shifts to Excel.")

    write_block(morning_shifts, ('E', 'F', 'G', 'H'), 'Morning', 7)
    write_block(afternoon_shifts, ('J', 'K', 'L', 'M'), 'Afternoon', 9)

    log_cb(f"Writing {len(leave_list)} flagged requests to column AG...")
    for i, entry in enumerate(leave_list[:10]): # Max 10 rows
        ws[f'AG{9+i}'] = entry

    log_cb(f"Writing {len(notes_list)} shift notes to column AG...")
    for i, note in enumerate(notes_list[:10]): # Max 10 rows
        ws[f'AG{19+i}'] = note

    if len(leave_list) > 10 or len(notes_list) > 10:
        log_cb("WARNING: only the first 10 requests and 10 notes fit in the template.")

    output_path = output_dir / f"{safe_date}.xlsm"
    log_cb(f"Saving workbook to: {output_path}")
    wb.save(output_path)
    return str(output_path)


# ==========================================
# 5. CUSTOM UI COMPONENTS
# ==========================================
class MiniCalendar(ctk.CTkToplevel):
    """Small themed date picker that writes directly to the target entry."""
    def __init__(self, parent, entry_widget):
        super().__init__(parent)
        self.entry_widget = entry_widget
        self.title("Select Date")
        self.geometry(f"300x330+{parent.winfo_rootx() + 150}+{parent.winfo_rooty() + 150}")
        self.configure(fg_color=CONFIG["bg_main"])
        self.resizable(False, False)
        self.curr_year, self.curr_month = datetime.now().year, datetime.now().month

        self.header_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.header_frame.pack(fill="x", padx=18, pady=(18, 10))
        ctk.CTkButton(self.header_frame, text="‹", width=36, height=32, corner_radius=12,
                      fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"],
                      command=self.prev_month).pack(side="left")
        self.month_lbl = ctk.CTkLabel(self.header_frame, text="", text_color=CONFIG["fg_main"],
                                      font=(CONFIG["font_family"], 16, "bold"))
        self.month_lbl.pack(side="left", expand=True)
        ctk.CTkButton(self.header_frame, text="›", width=36, height=32, corner_radius=12,
                      fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"],
                      command=self.next_month).pack(side="right")

        self.cal_frame = ctk.CTkFrame(self, fg_color=CONFIG["bg_surface"], corner_radius=18)
        self.cal_frame.pack(padx=18, pady=(0, 18), fill="both", expand=True)
        self.update_calendar()
        self.transient(parent)
        self.grab_set()

    def prev_month(self):
        self.curr_month = 12 if self.curr_month == 1 else self.curr_month - 1
        self.curr_year -= 1 if self.curr_month == 12 else 0
        self.update_calendar()

    def next_month(self):
        self.curr_month = 1 if self.curr_month == 12 else self.curr_month + 1
        self.curr_year += 1 if self.curr_month == 1 else 0
        self.update_calendar()

    def select_date(self, day):
        self.entry_widget.delete(0, tk.END)
        self.entry_widget.insert(0, f"{self.curr_month}-{day:02d}-{self.curr_year}")
        self.destroy()

    def update_calendar(self):
        for widget in self.cal_frame.winfo_children():
            widget.destroy()
        self.month_lbl.configure(text=f"{calendar.month_name[self.curr_month]} {self.curr_year}")
        for i, day_name in enumerate(["Sun", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]):
            ctk.CTkLabel(self.cal_frame, text=day_name, text_color=CONFIG["fg_muted"],
                         font=(CONFIG["font_family"], 12, "bold")).grid(row=0, column=i, padx=3, pady=(12, 5))
        for row_idx, week in enumerate(calendar.Calendar(firstweekday=calendar.SUNDAY).monthdayscalendar(self.curr_year, self.curr_month)):
            for col_idx, day in enumerate(week):
                if day:
                    ctk.CTkButton(self.cal_frame, text=str(day), width=29, height=27, corner_radius=10,
                                  fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["accent"],
                                  command=lambda selected=day: self.select_date(selected)).grid(
                                      row=row_idx + 1, column=col_idx, padx=3, pady=3)


# ==========================================
# 6. MAIN APPLICATION CLASS
# ==========================================
class SpreaderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("TPA Watch Desk Spreader Creator - v1.0.1")

        # Load the custom icon for the title bar
        try:
            self.root.iconbitmap(APP_ICON)
        except Exception:
            pass # Fails safely if the network drive drops or the icon is moved

        # Default to smaller window, hiding the log on startup
        self.root.geometry("760x690")
        self.root.minsize(700, 620)
        self.root.configure(fg_color=CONFIG["bg_main"])
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("dark-blue")
        self.build_ui()
        self.log("System Initialized. Awaiting user input.")

    def log(self, message):
        self.log_text.insert("end", f"{message}\n")
        self.log_text.see("end")
        self.root.update()

    def build_ui(self):
        main_container = ctk.CTkFrame(self.root, fg_color="transparent")
        main_container.pack(fill="both", expand=True, padx=30, pady=26)

        header = ctk.CTkFrame(main_container, fg_color="transparent")
        header.pack(fill="x", pady=(0, 20))
        ctk.CTkLabel(header, text="Daily Spreader", text_color=CONFIG["fg_main"],
                     font=(CONFIG["font_family"], 30, "bold")).pack(anchor="w")
        ctk.CTkLabel(header, text="Semi-automation for the TPA Watch Desk", text_color=CONFIG["accent"],
                     font=(CONFIG["font_family"], 15, "bold")).pack(anchor="w", pady=(2, 0))
        ctk.CTkLabel(header, text="Pull schedule data from WMT and build the Excel spreader in a few clicks.",
                     text_color=CONFIG["fg_muted"], font=(CONFIG["font_family"], 14)).pack(anchor="w", pady=(7, 0))

        step1 = ctk.CTkFrame(main_container, fg_color=CONFIG["bg_surface"], corner_radius=CONFIG["radius_card"])
        step1.pack(fill="x", pady=(0, 14))
        ctk.CTkLabel(step1, text="01  GET THE SCHEDULE DATA", text_color=CONFIG["accent"],
                     font=(CONFIG["font_family"], 14, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", padx=24, pady=(20, 4))
        ctk.CTkLabel(step1, text="Open the WMT Worksheet View, then save the target date as HTML.",
                     text_color=CONFIG["fg_muted"], font=(CONFIG["font_family"], 13)).grid(row=1, column=0, sticky="w", padx=24, pady=(0, 18))
        ctk.CTkButton(step1, text="Launch WMT Scheduler", command=self.launch_wmt, width=210, height=42,
                      corner_radius=CONFIG["radius_control"], fg_color=CONFIG["btn_secondary"],
                      hover_color=CONFIG["btn_sec_hover"], font=(CONFIG["font_family"], 13, "bold")).grid(row=1, column=1, padx=(10, 24), pady=(0, 18))

        step2 = ctk.CTkFrame(main_container, fg_color=CONFIG["bg_surface"], corner_radius=CONFIG["radius_card"])
        step2.pack(fill="x", pady=(0, 16))
        ctk.CTkLabel(step2, text="02  BUILD THE SPREADER", text_color=CONFIG["accent"],
                     font=(CONFIG["font_family"], 14, "bold")).grid(row=0, column=0, columnspan=3, sticky="w", padx=24, pady=(20, 15))

        ctk.CTkLabel(step2, text="Target date", text_color=CONFIG["fg_main"], font=(CONFIG["font_family"], 13, "bold")).grid(row=1, column=0, sticky="w", padx=(24, 12), pady=8)
        self.date_entry = ctk.CTkEntry(step2, height=40, corner_radius=CONFIG["radius_control"], fg_color=CONFIG["bg_input"],
                                       border_color=CONFIG["border"], text_color=CONFIG["fg_main"], font=(CONFIG["font_family"], 13))
        self.date_entry.grid(row=1, column=1, sticky="ew", padx=5, pady=8)
        tmrw = datetime.now() + timedelta(days=1)
        self.date_entry.insert(0, f"{tmrw.month}-{tmrw.strftime('%d-%Y')}")
        ctk.CTkButton(step2, text="Select date", command=lambda: MiniCalendar(self.root, self.date_entry), width=105, height=40,
                      corner_radius=CONFIG["radius_control"], fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"],
                      font=(CONFIG["font_family"], 12, "bold")).grid(row=1, column=2, padx=(8, 24), pady=8)

        ctk.CTkLabel(step2, text="Worksheet HTML file", text_color=CONFIG["fg_main"], font=(CONFIG["font_family"], 13, "bold")).grid(row=2, column=0, sticky="w", padx=(24, 12), pady=(8, 20))
        self.file_entry = ctk.CTkEntry(step2, height=40, corner_radius=CONFIG["radius_control"], fg_color=CONFIG["bg_input"],
                                       border_color=CONFIG["border"], text_color=CONFIG["fg_main"], font=(CONFIG["font_family"], 12))
        self.file_entry.grid(row=2, column=1, sticky="ew", padx=5, pady=(8, 20))
        ctk.CTkButton(step2, text="Browse…", command=self.browse_file, width=105, height=40, corner_radius=CONFIG["radius_control"],
                      fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"], font=(CONFIG["font_family"], 12, "bold")).grid(row=2, column=2, padx=(8, 24), pady=(8, 20))
        step2.grid_columnconfigure(1, weight=1)

        self.dry_run_var = tk.BooleanVar(value=False)
        ctk.CTkCheckBox(main_container, text="Preview / validate only (do not save Excel)",
                        variable=self.dry_run_var, font=(CONFIG["font_family"], 12)).pack(anchor="w", pady=(0, 8))

        self.btn_generate = ctk.CTkButton(main_container, text="Generate Spreader", command=self.run_generation, height=54,
                                          corner_radius=27, fg_color=CONFIG["accent"], hover_color=CONFIG["accent_hover"],
                                          text_color=CONFIG["fg_on_accent"], font=(CONFIG["font_family"], 16, "bold"))
        self.btn_generate.pack(fill="x", pady=(3, 14))

        actions = ctk.CTkFrame(main_container, fg_color="transparent")
        actions.pack(fill="x")
        self.btn_help = ctk.CTkButton(actions, text="Help & Instructions", command=self.show_help, height=34, corner_radius=17,
                                      fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"], font=(CONFIG["font_family"], 12, "bold"))
        self.btn_help.pack(side="left", padx=(0, 8))
        self.btn_toggle = ctk.CTkButton(actions, text="Show Debug Logger ↓", command=self.toggle_log, height=34, corner_radius=17,
                                        fg_color=CONFIG["btn_secondary"], hover_color=CONFIG["btn_sec_hover"], font=(CONFIG["font_family"], 12, "bold"))
        self.btn_toggle.pack(side="left")

        self.log_frame = ctk.CTkFrame(main_container, fg_color=CONFIG["bg_surface"], corner_radius=CONFIG["radius_card"])
        ctk.CTkLabel(self.log_frame, text="SYSTEM LOG", text_color=CONFIG["accent"], font=(CONFIG["font_family"], 13, "bold")).pack(anchor="w", padx=18, pady=(15, 8))
        self.log_text = ctk.CTkTextbox(self.log_frame, height=220, corner_radius=16, fg_color=CONFIG["bg_log"],
                                       text_color=CONFIG["fg_log"], font=("Consolas", 12), wrap="word")
        self.log_text.pack(fill="both", expand=True, padx=14, pady=(0, 14))

    def toggle_log(self):
        if self.log_frame.winfo_ismapped():
            self.log_frame.pack_forget()
            self.btn_toggle.configure(text="Show Debug Logger ↓")
            self.root.geometry("760x690")
        else:
            self.log_frame.pack(fill="both", expand=True, pady=(14, 0))
            self.btn_toggle.configure(text="Hide Debug Logger ↑")
            self.root.geometry("760x930")

    def show_help(self):
        help_win = ctk.CTkToplevel(self.root)
        help_win.title("Help & Instructions")
        help_win.geometry("550x650")
        help_win.configure(fg_color=CONFIG["bg_main"])

        try:
            help_win.iconbitmap(APP_ICON)
        except Exception:
            pass

        title = ctk.CTkLabel(help_win, text="Spreader Creator - User Guide", text_color=CONFIG["fg_main"], font=(CONFIG["font_family"], 20, "bold"))
        title.pack(pady=(15, 10))

        text_frame = ctk.CTkFrame(help_win, fg_color="transparent")
        text_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 20))
        help_text = ctk.CTkTextbox(text_frame, wrap="word", fg_color=CONFIG["bg_input"], text_color=CONFIG["fg_main"],
                                   font=(CONFIG["font_family"], 13), corner_radius=16)
        help_text.pack(fill="both", expand=True)

        instructions = (
            "1. DOWNLOADING WMT DATA\n"
            "----------------------------------------\n"
            "• Click 'Launch WMT Scheduler' to open Chrome.\n"
            "• Navigate to the Worksheet View for your target date.\n"
            "• Press Ctrl+S to save the page.\n"
            "• CRITICAL: Change 'Save as type' to 'Webpage, HTML Only'.\n"
            "• Save it to the network Downloads folder.\n\n"

            "2. GENERATING THE SPREADER\n"
            "----------------------------------------\n"
            "• Select the Target Date using the built-in calendar.\n"
            "• Click 'Browse' and select the HTML file you just saved.\n"
            "• Click 'Generate Spreader'.\n\n"

            "3. HOW IT WORKS BEHIND THE SCENES\n"
            "----------------------------------------\n"
            "• DEV Shifts: Automatically marked with a 'T' code. Their End Time is placed in the REQ column, and OFF is left blank so they don't count towards staffing numbers.\n"
            "• Overtime ($ and $$): Automatically recognized. The script knows standard 8-hour OT ($) vs extended 10-hour OT ($$).\n"
            "• FOS (Early Starts): Automatically adjusts the start time for the math, sets an 8-hour OFF time from that early start, and drops the final extended end time into the REQ column.\n"
            "• Leave Requests: Automatically flags unapproved leave (both full and partial shifts with times) to column AG. Ignores approved leave, shift changes, and shift swaps.\n"
            "• Shift Notes: Imports notes cleanly while filtering out automated WMT system notes (like '65 - OT OJT').\n\n"

            "4. TROUBLESHOOTING\n"
            "----------------------------------------\n"
            "• Error: 'No valid shifts found' -> The WMT file was likely saved as 'Webpage, Complete' instead of 'HTML Only'. Delete it, re-save correctly, and try again.\n"
            "• File won't open automatically -> The spreader was created successfully, but Excel might be busy. Navigate to the Watch Desk Documents folder to open it manually."
        )

        help_text.insert("end", instructions)
        help_text.configure(state="disabled")
        help_win.transient(self.root)

    def launch_wmt(self):
        try:
            paths = [r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"]
            chrome_path = next((p for p in paths if os.path.exists(p)), None)

            if chrome_path:
                webbrowser.register('chrome_explicit', None, webbrowser.BackgroundBrowser(chrome_path))
                webbrowser.get('chrome_explicit').open_new_tab(WMT_URL)
            else:
                webbrowser.open_new_tab(WMT_URL)

            messagebox.showinfo("Step 1", "WMT Scheduler should be opening in Chrome!\n\n1. Navigate to target date.\n2. Press 'Ctrl + S'.\n3. Change 'Save as type' to 'Webpage, HTML Only'.\n4. Save to network Downloads folder.")
        except Exception as e:
            messagebox.showerror("Error", f"Could not open browser:\n{str(e)}")

    def browse_file(self):
        filepath = filedialog.askopenfilename(initialdir=DOWNLOADS_DIR, title="Select WMT HTML File", filetypes=[("HTML Files", "*.html *.htm")])
        if filepath:
            self.file_entry.delete(0, tk.END)
            self.file_entry.insert(0, filepath)

    def run_generation(self):
        self.log("--- INITIATING SPREADER GENERATION ---")
        filepath = self.file_entry.get()
        if not filepath:
            self.log("ERROR: No HTML file path provided.")
            messagebox.showwarning("Missing File", "Please select a WMT HTML file first.")
            return

        try:
            parsed = parse_html_data(filepath, self.log)
            if not parsed.shifts:
                self.log("PROCESS HALTED: No valid shifts found.")
                messagebox.showerror("No Data", "Could not find any shifts. Did you save the WMT page correctly?")
                return

            if self.dry_run_var.get():
                preview = (
                    f"Shifts: {len(parsed.shifts)}\n"
                    f"Flagged requests: {len(parsed.leave_list)}\n"
                    f"Shift notes: {len(parsed.notes_list)}\n\n"
                    "First shifts:\n"
                    + "\n".join(
                        f"  {shift['initials']} | {shift['start']} | {shift['code'] or 'Standard'}"
                        for shift in parsed.shifts[:20]
                    )
                )
                build_excel_file(parsed.shifts, parsed.leave_list, parsed.notes_list,
                                 self.date_entry.get(), self.log, dry_run=True)
                messagebox.showinfo("Preview", preview)
                return

            out_path = build_excel_file(parsed.shifts, parsed.leave_list, parsed.notes_list,
                                        self.date_entry.get(), self.log)
            self.log("--- SUCCESS! SPREADER COMPLETE ---")

            if messagebox.askyesno("Success", f"Spreader created successfully!\n\nSaved as: {out_path}\n\nWould you like to open it now?"):
                try: os.startfile(out_path)
                except Exception as e:
                    self.log(f"Native open failed: {e}")
                    messagebox.showerror("Error", "Could not open automatically. File is in the Watch Desk folder.")

        except Exception as e:
            self.log(f"CRITICAL ERROR: {str(e)}")
            self.log(traceback.format_exc())
            messagebox.showerror("Error", "A critical error occurred. Check the debug log for details.")


if __name__ == "__main__":
    root = ctk.CTk()
    app = SpreaderApp(root)
    root.mainloop()
