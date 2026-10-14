from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

import openpyxl

from business_rules import calculate_times, sort_chronological, to_num
from models import ParsedData


def validate_date(date_str: str) -> str:
    parsed = datetime.strptime(date_str.strip(), "%m-%d-%Y")
    return f"{parsed.month}-{parsed.day:02d}-{parsed.year}"


def build_excel_file(data: ParsedData, date_str: str, template_path: str, output_dir: str,
                     log_cb: Callable[[str], None] = lambda _: None,
                     dry_run: bool = False) -> Optional[str]:
    safe_date = validate_date(date_str)
    if dry_run:
        log_cb("DRY RUN: validated input; no workbook was opened or saved.")
        return None
    template = Path(template_path)
    destination = Path(output_dir) / f"{safe_date}.xlsm"
    if not template.is_file():
        raise FileNotFoundError(f"Excel template not found: {template}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(template, keep_vba=True)
    if "Projected Coverage" not in wb.sheetnames:
        raise KeyError("Template is missing the 'Projected Coverage' worksheet")
    ws = wb["Projected Coverage"]

    morning = [s for s in data.shifts if s["start"] in {"2230", "2300", "2330", "2400"} or int(s["start"]) < 1130]
    afternoon = [s for s in data.shifts if s not in morning]
    morning.sort(key=sort_chronological)
    afternoon.sort(key=sort_chronological)

    # Clear the template's writable ranges so records from a prior run cannot survive.
    for row in list(range(7, 7 + 100)) + list(range(9, 9 + 100)):
        for col in ("E", "F", "G", "H", "J", "K", "L", "M"):
            ws[f"{col}{row}"] = None
    for row in list(range(9, 19)) + list(range(19, 29)):
        ws[f"AG{row}"] = None

    def write_block(shifts, cols, start_row):
        row, last_start = start_row, ""
        for shift in shifts:
            if last_start and shift["start"] != last_start:
                row += 1
            if row >= start_row + 100:
                raise ValueError("Too many shifts for the available template rows")
            off, req = calculate_times(shift["start"], shift["code"], shift["orig_code"], shift["ot"], shift["fos"])
            for col, value in zip(cols, (shift["initials"], to_num(off), shift["code"], to_num(req))):
                ws[f"{col}{row}"] = value
            last_start, row = shift["start"], row + 1

    write_block(morning, ("E", "F", "G", "H"), 7)
    write_block(afternoon, ("J", "K", "L", "M"), 9)
    if len(data.leave_list) > 10 or len(data.notes_list) > 10:
        log_cb("Warning: only the first 10 requests and 10 notes fit in the template.")
    for i, entry in enumerate(data.leave_list[:10]):
        ws[f"AG{9 + i}"] = entry
    for i, note in enumerate(data.notes_list[:10]):
        ws[f"AG{19 + i}"] = note
    wb.save(destination)
    return str(destination)
